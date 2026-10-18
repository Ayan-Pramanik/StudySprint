"""Format-aware conversion into editable Markdown and immutable image assets."""
from __future__ import annotations

import hashlib
import os
import re
from collections import Counter
from decimal import Decimal, InvalidOperation
from pathlib import Path
from uuid import uuid4

from PIL import Image

ENGINE = "docling-2.140.0/studysprint-1"
MAX_IMAGES = 200
MAX_IMAGE_BYTES = 8 * 1024 * 1024
MAX_IMAGE_TOTAL = 40 * 1024 * 1024


def _fail(message):
    from .conversion import ConversionError
    raise ConversionError(message)


def _pdf_options():
    from docling.datamodel.pipeline_options import PdfPipelineOptions, HeadingHierarchyOptions
    root = Path(os.environ.get("CONVERTER_MODEL_DIR", Path(__file__).parent.parent / ".sites-runtime/docling-models")).resolve()
    if not root.is_dir():
        _fail("The local layout models are missing. Run converter setup before uploading documents.")
    return PdfPipelineOptions(artifacts_path=root, do_ocr=False, generate_page_images=True, generate_picture_images=True,
        generate_parsed_pages=True, images_scale=2, document_timeout=500,
        heading_hierarchy_options=HeadingHierarchyOptions(enabled=True), enable_remote_services=False, allow_external_plugins=False)


def _converter():
    from docling.document_converter import DocumentConverter, PdfFormatOption, PowerpointFormatOption, WordFormatOption
    from docling.datamodel.backend_options import MsPowerpointBackendOptions, MsWordBackendOptions
    from docling.datamodel.base_models import InputFormat
    return DocumentConverter(allowed_formats=[InputFormat.PDF, InputFormat.PPTX, InputFormat.DOCX], format_options={
        InputFormat.PDF: PdfFormatOption(pipeline_options=_pdf_options()),
        InputFormat.PPTX: PowerpointFormatOption(backend_options=MsPowerpointBackendOptions(render_chart_images=False)),
        InputFormat.DOCX: WordFormatOption(backend_options=MsWordBackendOptions(render_chart_images=False)),
    })


def _successful(converter, source):
    from docling.datamodel.base_models import ConversionStatus
    result = converter.convert(source, raises_on_error=False)
    if result.status != ConversionStatus.SUCCESS:
        _fail("The layout converter could not process the complete document. Split it into smaller files or retry with a fresh copy.")
    return result.document


def _repair_pdf(doc):
    """Separate text blocks merged across disjoint columns using provenance spans."""
    from docling_core.types.doc import DocItemLabel, ContentLayer
    callouts = {}
    candidates = []
    for text in list(doc.texts):
        if len(text.prov) == 1:
            prov = text.prov[0]; size = doc.pages[prov.page_no].size
            box = prov.bbox.to_top_left_origin(size.height)
            if box.t < size.height * .12 and len(text.text) < 90:
                candidates.append((re.sub(r"\W", "", text.text).lower(), text))
        if text.label != DocItemLabel.TEXT or len(text.prov) < 2:
            continue
        first = text.prov[0]
        separate = [p for p in text.prov[1:] if p.page_no == first.page_no and (p.bbox.l > first.bbox.r + 12 or p.bbox.r < first.bbox.l - 12)]
        if not separate:
            continue
        original = text.text
        main = [original[p.charspan[0]:p.charspan[1]] for p in text.prov if p not in separate]
        text.text = text.orig = " ".join(main).strip()
        text.prov = [p for p in text.prov if p not in separate]
        parent = text.parent.resolve(doc) if text.parent else doc.body
        position = next((i for i, ref in enumerate(parent.children) if ref.cref == text.self_ref), len(parent.children))
        for index, prov in enumerate(separate):
            content = original[prov.charspan[0]:prov.charspan[1]].strip()
            marker = "SSCALLOUT" + uuid4().hex
            added = doc.add_text(label=DocItemLabel.TEXT, text=marker, parent=parent)
            added.prov = [prov]
            ref = parent.children.pop()
            parent.children.insert(position + index + 1, ref)
            callouts[marker] = content
    counts = Counter(key for key, _ in candidates)
    for key, text in candidates:
        if counts[key] >= 2:
            text.content_layer = ContentLayer.FURNITURE
    headings = [t for t in doc.texts if hasattr(t, "level") and t.content_layer == ContentLayer.BODY]
    if headings:
        minimum = headings[0].level
        for heading in headings: heading.level = max(1, heading.level - minimum + 1)
    return callouts


def _native_spacing(doc,preview):
    """Use native word spacing only when every letter/digit agrees with extraction."""
    import pypdfium2 as pdfium
    canonical=lambda text:"".join(c.casefold() for c in text if c.isalnum())
    with pdfium.PdfDocument(preview) as pdf:
        for item in doc.texts:
            if len(item.prov)!=1 or not item.text.strip():continue
            prov=item.prov[0];size=doc.pages[prov.page_no].size
            page=pdf[prov.page_no-1]
            # Docling's image coordinates need no transform on these pages.
            # Leave rotated/cropped geometry to the layout backend.
            if page.get_rotation() or abs(page.get_width()-size.width)>.5 or abs(page.get_height()-size.height)>.5:
                page.close();continue
            left,bottom,_,_=page.get_bbox()
            if abs(left)>.5 or abs(bottom)>.5:page.close();continue
            box=prov.bbox.to_bottom_left_origin(size.height);text_page=page.get_textpage()
            raw=text_page.get_text_bounded(box.l-1,box.b-1,box.r+1,box.t+1)
            text_page.close();page.close()
            raw=re.sub(r"([A-Za-z]{2,})-[ \t]*\r?\n[ \t]*([a-z]{2,})",r"\1\2",raw.replace("\ufffe",""))
            raw=re.sub(r"\s+"," ",raw).strip()
            if raw and canonical(raw)==canonical(item.text):item.text=item.orig=raw


def _crop(preview, doc, item):
    if not item.prov:
        return None
    prov = item.prov[0]
    if not prov.page_no or prov.page_no not in doc.pages:
        return None
    size = doc.pages[prov.page_no].size
    if not size or not size.width or not size.height:
        return None
    box = prov.bbox.to_top_left_origin(size.height)
    if box.r <= box.l or box.b <= box.t:
        return None
    import pypdfium2 as pdfium
    with pdfium.PdfDocument(preview) as pdf:
        page = pdf[prov.page_no - 1]
        bitmap = page.render(scale=2)
        full = bitmap.to_pil()
        bounds = (max(0, int(box.l / size.width * full.width)-8), max(0, int(box.t / size.height * full.height)-8), min(full.width, int(box.r / size.width * full.width + 1)+8), min(full.height, int(box.b / size.height * full.height + 1)+8))
        image = full.crop(bounds)
        full.close(); bitmap.close(); page.close()
        return image


def _chart_numbers(source, doc):
    from pptx import Presentation
    presentation = Presentation(source)
    formats = {}
    def walk(shapes):
        for shape in shapes:
            if hasattr(shape, "shapes"): yield from walk(shape.shapes)
            yield shape
    for page, slide in enumerate(presentation.slides, 1):
        for shape in walk(slide.shapes):
            if not shape.has_chart: continue
            code = None
            for plot in shape.chart.plots:
                if plot.has_data_labels:
                    code = plot.data_labels.number_format
                    if code and code != "General": break
            if not code or code == "General":
                try: code = shape.chart.value_axis.tick_labels.number_format
                except (ValueError, AttributeError): pass
            formats[page] = code or "General"
    for picture in doc.pictures:
        table = picture.meta.tabular_chart.chart_data if picture.meta and picture.meta.tabular_chart else None
        if not table: continue
        code = formats.get(picture.prov[0].page_no if picture.prov else 0, "General")
        precision = re.search(r"\.([0#]+)", code)
        currency = next((c for c in "$€£¥" if c in code), "")
        for cell in table.table_cells:
            if cell.column_header or cell.row_header or not cell.text: continue
            try:
                number = Decimal(cell.text)
                if code != "General":
                    section=code.split(";")[1 if number<0 and ";" in code else 0]
                    if "%" in section: number *= 100
                    decimals=min(10,len(precision[1])) if precision else 0
                    value=format(abs(number),("," if "," in section else "")+f".{decimals}f")
                    value=currency+value+("%" if "%" in section else "")
                    if number<0:value="("+value+")" if "(" in section and ")" in section else "-"+value
                    cell.text=value
                else:
                    cell.text = f"{float(number):.12g}"
            except (InvalidOperation, ValueError): pass
    return presentation


def extract_document(source: Path, preview: Path, work: Path):
    from docling_core.types.doc import ImageRef, ImageRefMode, ContentLayer
    converter = _converter()
    is_pdf = source.suffix.lower() == ".pdf"
    doc = _successful(converter, preview if is_pdf else source)
    callouts = _repair_pdf(doc) if is_pdf else {}
    if is_pdf:_native_spacing(doc,preview)
    presentation = _chart_numbers(source, doc) if source.suffix.lower() == ".pptx" else None
    assets, digest_ids, total, warnings, extras = [], {}, 0, [], []
    missing_native = []
    (work / "assets").mkdir(exist_ok=True)

    def save_image(image, item, context_doc):
        nonlocal total
        if image is None:
            _fail("A document figure could not be rendered. Save a fresh source copy and retry.")
        if image.width * image.height > 25_000_000:
            _fail("A document figure exceeds the image resolution limit.")
        image = image.convert("RGB") if image.mode not in {"RGB", "RGBA", "L"} else image
        import io
        buffer = io.BytesIO(); image.save(buffer, "PNG", optimize=True); data = buffer.getvalue()
        digest = hashlib.sha256(data).hexdigest()
        if digest in digest_ids:
            return digest_ids[digest]
        if len(assets) >= MAX_IMAGES or len(data) > MAX_IMAGE_BYTES or total + len(data) > MAX_IMAGE_TOTAL:
            _fail("The document exceeds the limit of 200 images, 8 MiB per image, or 40 MiB of images. Split it into smaller files.")
        asset_id = str(uuid4()); total += len(data); digest_ids[digest] = asset_id
        (work / "assets" / f"{asset_id}.png").write_bytes(data)
        asset = {"id": asset_id, "mediaType": "image/png", "size": len(data), "sha256": digest, "width": image.width, "height": image.height}
        if item.prov:
            prov = item.prov[0]; asset.update(page=prov.page_no, bbox=[prov.bbox.l, prov.bbox.t, prov.bbox.r, prov.bbox.b])
        if hasattr(item, "caption_text"):
            caption=item.caption_text(context_doc)
            if not caption and item.prov:
                prov=item.prov[0]
                nearby=[t for t in context_doc.texts if str(t.label.value)=="caption" and t.prov and t.prov[0].page_no==prov.page_no]
                if nearby:
                    def distance(t):
                        a,b=prov.bbox,t.prov[0].bbox
                        return max(a.l-b.r,b.l-a.r,0)**2+max(min(a.t,a.b)-max(b.t,b.b),min(b.t,b.b)-max(a.t,a.b),0)**2
                    caption=min(nearby,key=distance).text
            asset["alt"] = caption[:1000]
        assets.append(asset)
        return asset_id

    for picture in doc.pictures:
        image = _crop(preview, doc, picture) if is_pdf or presentation else None
        if image is None:
            image = picture.get_image(doc)
        if image is None and source.suffix.lower() == ".docx":
            # Native Word objects without image bytes are recovered below from
            # the rendered preview, where a trustworthy page location exists.
            warnings.append("An unsupported Word visual is included under Additional figures.")
            missing_native.append(picture)
            continue
        asset_id = save_image(image, picture, doc)
        picture.image = ImageRef.from_pil(image, dpi=144)
        picture.image.uri = Path(f"asset-{asset_id}.png")
        image.close()

    # Word's native backend has semantic styles but cannot locate every floating
    # drawing. Compare preview figures to the native pictures using thumbnails.
    if source.suffix.lower() == ".docx":
        preview_doc = _successful(converter, preview)
        native_thumbs = []
        for asset in assets:
            with Image.open(work / "assets" / f'{asset["id"]}.png') as image:
                native_thumbs.append((image.width / image.height, list(image.convert("L").resize((16,16)).get_flattened_data())))
        for picture in preview_doc.pictures:
            image = _crop(preview, preview_doc, picture)
            if image is None: continue
            thumb = list(image.convert("L").resize((16,16)).get_flattened_data())
            similar = any(abs(image.width / image.height - ratio) < .15 and sum(abs(a-b) for a,b in zip(thumb, other))/256 < 24 for ratio, other in native_thumbs)
            if similar: image.close(); continue
            asset_id = save_image(image, picture, preview_doc); image.close()
            caption = picture.caption_text(preview_doc)
            page = picture.prov[0].page_no if picture.prov else 1
            extras.append(f'### Figure from page {page}\n\n![{caption or "Document figure"}](asset:{asset_id})'+(f'\n\n{caption}' if caption else ''))
        # Native Word extraction cannot represent VML, SmartArt, embedded
        # objects and some vector shapes. Without reliable object-to-page
        # anchors, retain labelled page crops rather than omitting the visual.
        import zipfile
        with zipfile.ZipFile(source) as package:
            xml=package.read("word/document.xml")
        unsupported=any(token in xml for token in [b"<w:pict",b"<w:object",b"<wps:wsp",b"<dgm:relIds",b"<c:chart",b"<v:shape"])
        if unsupported:
            import pypdfium2 as pdfium
            from types import SimpleNamespace
            with pdfium.PdfDocument(preview) as pdf:
                for number,page in enumerate(pdf,1):
                    bitmap=page.render(scale=1.5);image=bitmap.to_pil()
                    asset_id=save_image(image,SimpleNamespace(prov=[]),doc)
                    image.close();bitmap.close();page.close()
                    extras.append(f'### Original layout — page {number}\n\n![Unsupported Word visual, page {number}](asset:{asset_id})')
            warnings.append("Unsupported Word visual objects are retained in labelled page previews because their exact paragraph positions are uncertain.")
        if missing_native:
            if not extras: _fail("A Word visual could not be recovered from its PDF preview.")
            for item in missing_native:
                if item.parent:
                    parent=item.parent.resolve(doc)
                    parent.children=[ref for ref in parent.children if ref.cref != item.self_ref]
        if extras: warnings.append("Some floating Word figures have uncertain placement and are retained in a labelled Additional figures section.")

    # Retain faithful images for tables whose merged cells cannot be expressed
    # by a plain Markdown grid, alongside an editable linear cell description.
    complex_tables = {}
    for table in doc.tables:
        if not any(cell.row_span > 1 or cell.col_span > 1 for cell in table.data.table_cells): continue
        image = _crop(preview, doc, table)
        if image is None:
            warnings.append("A merged table is represented as editable text; refer to the Original view for its layout.")
            continue
        asset_id = save_image(image, table, doc); image.close()
        linear = "\n".join(f'- Row {cell.start_row_offset_idx+1}, column {cell.start_col_offset_idx+1}: {cell.text}' for cell in table.data.table_cells if cell.text)
        complex_tables[table.export_to_markdown(doc=doc)] = f'![Table layout](asset:{asset_id})\n\n**Table text**\n\n{linear}'

    if is_pdf:
        import pypdfium2 as pdfium
        repairs = {}
        with pdfium.PdfDocument(preview) as pdf:
            for page in pdf:
                text_page=page.get_textpage()
                raw=text_page.get_text_range()
                for match in re.finditer(r"([A-Za-z]{2,})(?:-[ \t]*\r?\n[ \t]*|\ufffe)([a-z]{2,})",raw):
                    repairs[match[1]+"- "+match[2]]=match[1]+match[2]
                text_page.close();page.close()
        for item in doc.texts:
            for old,new in repairs.items(): item.text=item.text.replace(old,new)
        for table in doc.tables:
            for cell in table.data.table_cells:
                for old,new in repairs.items(): cell.text=cell.text.replace(old,new)

    def export(page=None):
        return doc.export_to_markdown(image_mode=ImageRefMode.REFERENCED, page_no=page, included_content_layers={ContentLayer.BODY}, include_picture_classification=False)
    if presentation:
        chunks = []
        for index, slide in enumerate(presentation.slides, 1):
            title = slide.shapes.title.text.strip() if slide.shapes.title else ""
            body = export(index)
            if title: body = re.sub(r"^#{1,6}\s+" + re.escape(title) + r"\s*$", "", body, count=1, flags=re.MULTILINE)
            body = re.sub(r"^(#{1,6})\s", lambda m: "#" * min(6, len(m[1])+2) + " ", body, flags=re.MULTILINE)
            section = f'## Slide {index}'+(f' — {title}' if title else '')+f'\n\n{body.strip()}'
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes: section += f'\n\n### Speaker notes\n\n{notes}'
            chunks.append(section)
        markdown = "\n\n---\n\n".join(chunks)
    else:
        markdown = export()
    for marker, content in callouts.items(): markdown = markdown.replace(marker, "> " + content.replace("\n", "\n> "))
    for original, replacement in complex_tables.items(): markdown = markdown.replace(original, replacement)
    for asset in assets: markdown = markdown.replace(f'asset-{asset["id"]}.png', f'asset:{asset["id"]}')
    for asset in assets:
        if asset.get("alt"):
            alt=asset["alt"][:200].replace("\n"," ").replace("[",r"\[").replace("]",r"\]")
            markdown=markdown.replace(f'![Image](asset:{asset["id"]})',f'![{alt}](asset:{asset["id"]})')
    if extras: markdown += "\n\n## Additional figures\n\n" + "\n\n".join(extras)
    if is_pdf:
        # The legacy Markdown serializer emits every PDF heading as ## despite
        # the hierarchy levels in the document model. Apply levels explicitly.
        for heading in [t for t in doc.texts if hasattr(t,"level") and t.content_layer == ContentLayer.BODY]:
            markdown=re.sub(r"^#{1,6}\s+"+re.escape(heading.text)+r"[ \t]*$",lambda _:"#"*min(6,max(1,heading.level))+" "+heading.text,markdown,count=1,flags=re.MULTILINE)
    if "<!-- image -->" in markdown:
        _fail("A document image was not extracted completely. Save a fresh source copy and retry.")
    return markdown, assets, {"engine": ENGINE, "warnings": list(dict.fromkeys(warnings))}
