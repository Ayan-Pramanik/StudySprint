import hashlib
import shutil
import zipfile
import subprocess
from pathlib import Path

import pytest

from converter.conversion import ConversionError, MAX_MARKDOWN_BYTES, clean_markdown, convert, validate_input


def test_markdown_preserves_frontmatter_and_utf8(tmp_path):
    source = tmp_path / "source.md"
    original = Path(__file__).with_name("fixtures").joinpath("lecture.md").read_bytes()
    source.write_bytes(original)
    result = convert(source, tmp_path)
    assert result["markdown"] == original.decode("utf-8")
    assert result["preview"] is True
    assert result["base"]["kind"] == "text-pdf"


def test_plain_text_is_not_interpreted_or_reformatted(tmp_path):
    source = tmp_path / "source.txt"
    original = "  Heading\r\n\r\n$E=mc²$\n[[Memory]]\n"
    source.write_bytes(original.encode("utf-8"))
    assert convert(source, tmp_path)["markdown"] == original.replace("[[Memory]]", "\\[\\[Memory\\]\\]")


@pytest.mark.parametrize("filename,data", [
    ("source.pdf", b"actually text"),
    ("source.docx", b"PK not a zip"),
    ("source.pptx", b"not a presentation"),
    ("source.txt", b"\xff\xfe\x00\x00"),
    ("source.md", b"text\x00binary"),
])
def test_disguised_or_binary_inputs_fail(tmp_path, filename, data):
    source = tmp_path / filename
    source.write_bytes(data)
    with pytest.raises(ConversionError):
        validate_input(source)


def test_zip_bomb_and_paths_are_rejected(tmp_path):
    source = tmp_path / "source.docx"
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("../outside.xml", "unsafe")
    with pytest.raises(ConversionError, match="unsafe"):
        validate_input(source)


def test_pdf_page_limit_rejects_before_ocr(pdf_factory):
    with pytest.raises(ConversionError, match="100 pages"):
        validate_input(pdf_factory(pages=101))


def test_markdown_images_are_preserved_and_size_is_bounded():
    assert clean_markdown("# Topic\n![Diagram](Picture1.jpg)\n<img src='missing.png'>\nText") == "# Topic\n![Diagram](Picture1.jpg)\n<img src='missing.png'>\nText"
    with pytest.raises(ConversionError, match="1 MiB"):
        clean_markdown("é" * (MAX_MARKDOWN_BYTES // 2 + 1))


@pytest.mark.parametrize("message", [
    "[tesseract] took too long to OCR - skipping",
    "page too big, skipping OCR (32.0 MPixels > 25.0 MPixels --skip-big)",
])
def test_success_exit_with_missing_ocr_page_is_rejected(tmp_path, monkeypatch, message):
    # These are OCRmyPDF's actual successful-but-incomplete page warning forms.
    # Its zero exit code must not authorize a note missing part of the source.
    from converter import conversion

    monkeypatch.setattr(conversion, "_run", lambda arguments, cwd: subprocess.CompletedProcess(arguments, 0, "", message))
    with pytest.raises(ConversionError, match="every page"):
        conversion.ocr_pdf(tmp_path / "source.pdf", tmp_path / "searchable.pdf", tmp_path, "eng")


@pytest.mark.integration
@pytest.mark.parametrize("kind", ["digital", "scanned", "mixed"])
def test_real_pdf_text_and_ocr_keep_source_unchanged(tmp_path, pdf_factory, kind):
    if not shutil.which("tesseract"):
        pytest.skip("Tesseract is not installed")
    source = pdf_factory(kind)
    before = hashlib.sha256(source.read_bytes()).digest()
    result = convert(source, tmp_path)
    text = result["markdown"].lower()
    if kind in {"digital", "mixed"}:
        assert "digital memory lecture" in text
        assert "retrieval practice" in text
    if kind in {"scanned", "mixed"}:
        assert "scanned study material" in text
        # Layout extraction may represent aligned OCR words as a Markdown table.
        semantic_text = " ".join(text.replace("|", " ").split())
        assert "spaced repetition improves retention." in semantic_text
    assert result["ocr"] is True
    assert result["preview"] is True
    assert hashlib.sha256(source.read_bytes()).digest() == before


@pytest.mark.integration
def test_real_powerpoint_markdown_and_pdf_preview(tmp_path):
    from pptx import Presentation

    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("LibreOffice is not installed")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Retrieval practice"
    slide.placeholders[1].text = "Spaced repetition improves retention."
    source = tmp_path / "source.pptx"
    presentation.save(source)
    before = source.read_bytes()
    result = convert(source, tmp_path)
    assert "Retrieval practice" in result["markdown"]
    assert "Spaced repetition" in result["markdown"]
    assert result["preview"] is True
    assert (tmp_path / "preview.pdf").read_bytes().startswith(b"%PDF-")
    assert source.read_bytes() == before


@pytest.mark.integration
def test_real_word_markdown_and_pdf_preview(tmp_path):
    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("LibreOffice is not installed")
    source = tmp_path / "source.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr("word/document.xml", '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Memory and learning</w:t></w:r></w:p><w:p><w:r><w:t>Retrieval practice improves retention.</w:t></w:r></w:p></w:body></w:document>')
    before = source.read_bytes()
    result = convert(source, tmp_path)
    assert "Memory and learning" in result["markdown"]
    assert "Retrieval practice improves retention." in result["markdown"]
    assert result["preview"] is True
    assert (tmp_path / "preview.pdf").read_bytes().startswith(b"%PDF-")
    assert source.read_bytes() == before

@pytest.mark.integration
def test_word_native_styles_nested_lists_tables_images_and_floating_fallback(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    from PIL import Image
    from docx.shared import Inches
    document=Document();document.add_heading('Unicode αβ 中文',0);document.add_heading('Structure',1)
    paragraph=document.add_paragraph('Keep ');paragraph.add_run('bold').bold=True;paragraph.add_run(' and ');paragraph.add_run('italic').italic=True
    document.add_paragraph('First item',style='List Bullet');document.add_paragraph('Nested item',style='List Bullet 2')
    table=document.add_table(rows=2,cols=2);table.cell(0,0).text='Term';table.cell(0,1).text='Meaning';table.cell(1,0).text='Recall';table.cell(1,1).text='Practice'
    image=tmp_path/'figure.png';Image.new('RGB',(180,90),'blue').save(image);document.add_picture(str(image),width=Inches(2))
    paragraph=document.add_paragraph();paragraph._p.append(parse_xml('<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:v="urn:schemas-microsoft-com:vml"><v:rect style="width:100pt;height:60pt" fillcolor="#00ff00"><v:textbox><w:txbxContent><w:p><w:r><w:t>Vector study diagram</w:t></w:r></w:p></w:txbxContent></v:textbox></v:rect></w:pict>'))
    source=tmp_path/'source.docx';document.save(source);original=source.read_bytes();result=convert(source,tmp_path)
    assert 'Structure' in result['markdown']
    assert '**bold**' in result['markdown'] and '*italic*' in result['markdown']
    assert 'First item' in result['markdown'] and 'Nested item' in result['markdown']
    assert '| Term' in result['markdown']
    assert len(result['assets'])>=2
    assert 'Original layout — page 1' in result['markdown']
    assert result['conversion']['warnings']
    assert '<!-- image -->' not in result['markdown']
    assert original==source.read_bytes()


def test_single_markdown_reports_missing_local_images_without_broken_links(tmp_path):
    source=tmp_path/'source.md';source.write_text('# Study\n\n![Figure](images/figure.png)\n')
    result=convert(source,tmp_path)
    assert result['markdown']==source.read_text()
    assert 'images/figure.png' in result['conversion']['warnings'][0]
    assert 'folder or ZIP' in result['conversion']['warnings'][0]
    assert source.read_text().endswith('![Figure](images/figure.png)\n')
