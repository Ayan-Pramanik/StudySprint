from uuid import uuid4
import shutil

import pytest

from fastapi.testclient import TestClient

from converter.app import Settings, create_app


def client_for(tmp_path):
    return TestClient(create_app(Settings(secret="test-secret", artifacts=tmp_path)))


def test_health_does_not_reveal_secret_and_api_requires_auth(tmp_path):
    with client_for(tmp_path) as client:
        health = client.get("/health")
        assert health.status_code == 200
        assert "test-secret" not in health.text
        response = client.post(f"/convert/{uuid4()}", files={"file": ("lecture.md", b"# Topic")})
        assert response.status_code == 401


def test_text_upload_roundtrip_and_cleanup(tmp_path):
    with client_for(tmp_path) as client:
        job = uuid4()
        headers = {"Authorization": "Bearer test-secret"}
        response = client.post(f"/convert/{job}", headers=headers, files={"file": ("lecture.md", b"---\ntitle: Topic\n---\n# Lecture\n")})
        assert response.status_code == 200, response.text
        assert response.json()["markdown"] == "---\ntitle: Topic\n---\n# Lecture\n"
        assert response.json()["base"]["kind"] == "text-pdf"
        assert client.get(f"/artifacts/{job}/preview", headers=headers).status_code == 200
        assert client.delete(f"/artifacts/{job}", headers=headers).status_code == 204
        assert not list(tmp_path.iterdir())


def test_busy_service_reports_retry_after(tmp_path):
    app = create_app(Settings(secret="test-secret", artifacts=tmp_path))
    with TestClient(app) as client:
        app.state.conversion_lock.acquire()
        try:
            response = client.post(f"/convert/{uuid4()}", headers={"Authorization": "Bearer test-secret"}, files={"file": ("lecture.md", b"# Topic")})
            assert response.status_code == 503
            assert response.headers["retry-after"] == "3"
        finally:
            app.state.conversion_lock.release()


def test_upload_over_limit_is_rejected_before_conversion(tmp_path):
    with client_for(tmp_path) as client:
        response = client.post(f"/convert/{uuid4()}", headers={"Authorization": "Bearer test-secret"}, files={"file": ("lecture.txt", b"x" * (10 * 1024 * 1024 + 1))})
        assert response.status_code == 413
        assert not list(tmp_path.iterdir())


def test_invalid_uuid_and_wrong_file_type_are_rejected(tmp_path):
    with client_for(tmp_path) as client:
        headers = {"Authorization": "Bearer test-secret"}
        assert client.post("/convert/not-a-uuid", headers=headers, files={"file": ("lecture.md", b"# Topic")}).status_code == 422
        assert client.post(f"/convert/{uuid4()}", headers=headers, files={"file": ("lecture.exe", b"binary")}).status_code == 415


def test_worker_deadline_kills_process_and_cleans_files(tmp_path):
    app = create_app(Settings(secret="test-secret", artifacts=tmp_path, deadline=0.001))
    with TestClient(app) as client:
        response = client.post(f"/convert/{uuid4()}", headers={"Authorization": "Bearer test-secret"}, files={"file": ("lecture.txt", b"Text for a cancelled conversion")})
        assert response.status_code == 422
        assert "limit" in response.json()["detail"]
        assert not list(tmp_path.iterdir())
        # A timed-out job must release the single-flight slot.
        assert app.state.conversion_lock.acquire(blocking=False)
        app.state.conversion_lock.release()


@pytest.mark.integration
def test_pdf_ocr_through_isolated_api_worker(tmp_path, pdf_factory):
    if not shutil.which("tesseract"):
        pytest.skip("Tesseract is not installed")
    source = pdf_factory("scanned")
    with client_for(tmp_path / "artifacts") as client:
        response = client.post(f"/convert/{uuid4()}", headers={"Authorization": "Bearer test-secret"}, files={"file": ("lecture.pdf", source.read_bytes(), "application/pdf")})
        assert response.status_code == 200, response.text
        assert "SCANNED study material" in response.json()["markdown"]
        assert response.json()["ocr"] is True


@pytest.mark.integration
def test_office_preview_download_and_deletion_through_api(tmp_path):
    from pptx import Presentation
    import io

    if not (shutil.which("soffice") or shutil.which("libreoffice")):
        pytest.skip("LibreOffice is not installed")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "StudySprint preview fixture"
    slide.placeholders[1].text = "Recall strengthens memory."
    stream = io.BytesIO()
    presentation.save(stream)
    with client_for(tmp_path) as client:
        job = uuid4()
        headers = {"Authorization": "Bearer test-secret"}
        response = client.post(f"/convert/{job}", headers=headers, files={"file": ("lecture.pptx", stream.getvalue())})
        assert response.status_code == 200, response.text
        assert response.json()["preview"] is True
        assert "Recall strengthens memory." in response.json()["markdown"]
        preview = client.get(f"/artifacts/{job}/preview", headers=headers)
        assert preview.status_code == 200
        assert preview.headers["content-type"] == "application/pdf"
        assert preview.content.startswith(b"%PDF-")
        assert client.delete(f"/artifacts/{job}", headers=headers).status_code == 204
        assert client.get(f"/artifacts/{job}/preview", headers=headers).status_code == 404
