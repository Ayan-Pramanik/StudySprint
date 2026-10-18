#!/usr/bin/env python3
"""Generate synthetic, non-personal fixtures into an explicitly chosen directory.

Run with converter/.venv/bin/python scripts/make-document-fixtures.py /tmp/studysprint-fixtures
"""

from __future__ import annotations

import argparse
from pathlib import Path
import zipfile

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


def scan_image() -> Image.Image:
    font_path = next((path for path in [
        Path("/System/Library/Fonts/Supplemental/Arial.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
    ] if path.is_file()), None)
    font = ImageFont.truetype(str(font_path), 52) if font_path else ImageFont.load_default(size=52)
    image = Image.new("RGB", (1600, 700), "white")
    draw = ImageDraw.Draw(image)
    draw.text((80, 120), "SCANNED study material", font=font, fill="black")
    draw.text((80, 240), "Spaced repetition improves retention.", font=font, fill="black")
    return image


def write_pdf(path: Path, *, kind: str, pages: int = 1) -> None:
    pdf = canvas.Canvas(str(path), pagesize=(612, 792), invariant=1)
    image = scan_image() if kind in {"scanned", "mixed"} else None
    for index in range(pages):
        if kind in {"digital", "mixed"}:
            pdf.setFont("Helvetica", 18)
            pdf.drawString(40, 740, f"DIGITAL memory lecture {index + 1}")
            pdf.setFont("Helvetica", 12)
            pdf.drawString(40, 710, "Retrieval practice strengthens memory.")
        if image is not None:
            pdf.drawImage(ImageReader(image), 40, 300, width=530, height=232)
        pdf.showPage()
    pdf.save()


def write_docx(path: Path) -> None:
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        archive.writestr("_rels/.rels", '<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        archive.writestr("word/document.xml", '<?xml version="1.0"?><w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>Memory and learning</w:t></w:r></w:p><w:p><w:r><w:t>Retrieval practice improves retention.</w:t></w:r></w:p><w:p><w:r><w:t>This is a synthetic StudySprint fixture.</w:t></w:r></w:p></w:body></w:document>')


def generate(output: Path) -> None:
    output.mkdir(parents=True, exist_ok=True)
    for kind in ("digital", "scanned", "mixed"):
        write_pdf(output / f"{kind}.pdf", kind=kind)
    # Long enough to observe a processing job and test cancellation publication fences.
    write_pdf(output / "slow-scanned.pdf", kind="scanned", pages=12)
    write_docx(output / "lecture.docx")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Retrieval practice"
    slide.placeholders[1].text = "Spaced repetition improves retention.\nThis is a synthetic StudySprint fixture."
    presentation.save(output / "slides.pptx")
    (output / "note.md").write_bytes(b"---\ntitle: Synthetic Markdown lecture\ntags: [memory, testing]\n---\n\n# Markdown memory lecture\n\nRetrieval practice improves retention.\n\nRelated: [[Memory]].\n")
    (output / "note.txt").write_bytes("Plain text memory lecture\r\n\r\nRetrieval practice improves retention.\r\nWorking memory is limited.\nUnicode: café and déjà vu.\n".encode("utf-8"))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("output_directory", type=Path, help="Explicit fixture destination; the known synthetic fixture filenames are replaced.")
    arguments = parser.parse_args()
    generate(arguments.output_directory.resolve())
    print(f"Generated 7 document fixtures and a slow cancellation PDF in {arguments.output_directory.resolve()}")
